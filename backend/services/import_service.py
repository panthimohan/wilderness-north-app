"""User-file detection, column mapping, validation, and domain normalization.

Uploaded bytes stay in memory. XLSX files are read with formulas disabled and
are never written, evaluated, or extracted to arbitrary filesystem paths.
"""
from __future__ import annotations

import base64
import csv
import io
import math
import re
import secrets
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from backend.models import FlightCapacity, HouseholdOrder, OrderItem

MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_FILES = 5

ORDER_FIELDS = {
    "order_id": "Order ID",
    "quantity": "Item quantity (optional; defaults to 1)",
    "household_id": "Household ID",
    "batch_id": "Batch ID",
    "order_date": "Order date",
    "destination_community": "Destination community",
    "product_id": "Product ID",
    "product_name": "Product name",
    "weight_lb": "Item weight (lb)",
    "length_in": "Item length (in)",
    "width_in": "Item width (in)",
    "height_in": "Item height (in)",
}
FLIGHT_FIELDS = {
    "departure_id": "Flight / departure ID",
    "departure_date": "Departure date",
    "available_totes": "Tote capacity",
    "available_payload_lb": "Payload capacity (lb)",
    "available_volume_cuft": "Cargo volume capacity (ft³)",
    "flight_cost": "Flight cost (optional)",
    "currency": "Currency (optional; defaults to CAD when cost is provided)",
}
ORDER_REQUIRED = {
    "order_id", "household_id", "order_date", "destination_community",
    "weight_lb", "length_in", "width_in", "height_in",
}
FLIGHT_REQUIRED = {
    "departure_id", "departure_date", "available_totes", "available_payload_lb",
}

# Aliases are compared after lowercasing and removing punctuation/whitespace.
ALIASES: dict[str, set[str]] = {
    "quantity": {"quantity", "qty", "itemquantity", "unitquantity", "units", "count"},
    "order_id": {"orderid", "ordernumber", "orderno", "order", "ord", "transactionid", "purchaseid"},
    "household_id": {"householdid", "household", "householdnumber", "customerid", "customerno", "customer", "accountid"},
    "batch_id": {"batchid", "batch", "batchnumber", "batchno"},
    "order_date": {"orderdate", "date", "dateordered", "purchasedate", "createddate", "transactiondate"},
    "destination_community": {"destinationcommunity", "destination", "community", "town", "shipto", "shiptocommunity", "location"},
    "product_id": {"productid", "itemid", "sku", "productcode", "itemcode", "upc"},
    "product_name": {"productname", "item", "itemname", "description", "product", "productdescription"},
    "weight_lb": {"weightlb", "weight", "itemweight", "itemweightlb", "productweight", "weightpounds"},
    "length_in": {"lengthin", "length", "boxlength", "itemlength", "lengthinch", "lengthinches"},
    "width_in": {"widthin", "width", "boxwidth", "itemwidth", "widthinch", "widthinches"},
    "height_in": {"heightin", "height", "boxheight", "itemheight", "heightinch", "heightinches"},
    "departure_id": {"departureid", "flightid", "flight", "flightnumber", "flightno", "departure", "departurenumber"},
    "departure_date": {"departuredate", "flightdate", "departdate", "date"},
    "available_totes": {"availabletotes", "totecapacity", "maxtotes", "maximumtotes", "totes", "totelimit"},
    "available_payload_lb": {"availablepayloadlb", "payload", "payloadlb", "payloadcapacity", "weightcapacity", "weightcapacitylb", "maxpayload", "maxweight", "cargoweight"},
    "available_volume_cuft": {"availablevolumecuft", "volume", "volumecapacity", "cargovolume", "cargovolumecuft", "maxvolume", "volumeft3", "volumecuft"},
    "flight_cost": {"flightcost", "cost"},
    "currency": {"currency", "currencycode"},
}


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


@dataclass(slots=True)
class ImportDraft:
    draft_id: str
    filename: str
    headers: list[str] = field(default_factory=list)
    rows: list[dict[str, str]] = field(default_factory=list)
    dataset_type: str = "unknown"
    confidence: str = "low"
    signals: list[str] = field(default_factory=list)
    parse_error: str | None = None
    note: str | None = None


def safe_filename(filename: str) -> str:
    name = Path(filename.replace("\\", "/")).name.strip()
    if not name or name in {".", ".."} or any(ord(char) < 32 for char in name):
        raise ValueError("filename is missing or invalid")
    if Path(name).suffix.lower() not in {".csv", ".xlsx"}:
        raise ValueError("only .csv and .xlsx files are supported; macro-enabled workbooks are not accepted")
    return name


def decode_upload(filename: str, encoded_content: str) -> tuple[str, bytes]:
    name = safe_filename(filename)
    try:
        content = base64.b64decode(encoded_content, validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{name}: upload content is not valid base64") from exc
    if not content:
        raise ValueError(f"{name}: uploaded file is empty")
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise ValueError(f"{name}: file exceeds the {MAX_IMPORT_FILE_BYTES // (1024 * 1024)} MB limit")
    return name, content


def _parse_csv(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must use UTF-8 encoding") from exc
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if sample else csv.excel
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text, newline=""), dialect=dialect)
    if reader.fieldnames is None:
        raise ValueError("CSV has no header row")
    headers = [str(header or "").strip() for header in reader.fieldnames]
    if len(headers) > 100:
        raise ValueError("CSV exceeds the 100 column limit")
    if not headers or any(not header for header in headers):
        raise ValueError("CSV contains an empty column name")
    if len({_norm(header) for header in headers}) != len(headers):
        raise ValueError("CSV contains duplicate column names")
    rows: list[dict[str, str]] = []
    try:
        for number, row in enumerate(reader, start=2):
            if number > 10002:
                raise ValueError("CSV exceeds the 10,000 data row limit")
            if None in row:
                raise ValueError(f"row {number} has extra values beyond the header columns")
            rows.append({headers[index]: _cell_text(row.get(raw_header)) for index, raw_header in enumerate(reader.fieldnames)})
    except csv.Error as exc:
        raise ValueError(f"CSV could not be parsed: {exc}") from exc
    return headers, rows


def _parse_xlsx(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]], str | None]:
    if not content.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValueError("file extension is .xlsx but the content is not a valid XLSX workbook")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX support requires openpyxl; install project requirements") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("XLSX expands beyond the 50 MB processing limit")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.active
        if sheet.max_column and sheet.max_column > 100:
            raise ValueError("workbook exceeds the 100 column limit")
        if sheet.max_row and sheet.max_row > 10002:
            raise ValueError("workbook exceeds the 10,000 data row limit")
        raw_rows = sheet.iter_rows(values_only=True)
        header_row = next(raw_rows, None)
        if not header_row:
            raise ValueError("workbook has no header row")
        headers = [_cell_text(cell) for cell in header_row]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not header for header in headers):
            raise ValueError("workbook contains an empty column name")
        if len({_norm(header) for header in headers}) != len(headers):
            raise ValueError("workbook contains duplicate column names")
        rows: list[dict[str, str]] = []
        for row_number, values in enumerate(raw_rows, start=2):
            if row_number > 10002:
                raise ValueError("workbook exceeds the 10,000 data row limit")
            values = tuple(values)
            if len(values) > len(headers) and any(value not in (None, "") for value in values[len(headers):]):
                raise ValueError(f"row {row_number} has values beyond the header columns")
            if not any(value not in (None, "") for value in values):
                continue
            padded = values[:len(headers)] + (None,) * max(0, len(headers) - len(values))
            rows.append({headers[index]: _cell_text(padded[index]) for index in range(len(headers))})
        note = None
        if len(workbook.sheetnames) > 1:
            note = f"Only the active worksheet ({sheet.title}) was analyzed; this workbook has {len(workbook.sheetnames)} sheets."
        workbook.close()
        return headers, rows, note
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"could not read XLSX workbook: {exc}") from exc


def _has_incompatible_unit(header: str, field_name: str) -> bool:
    name = _norm(header)
    suffixes = {
        "weight_lb": ("kilograms", "kilogram", "kg", "grams", "gram", "g"),
        "available_payload_lb": ("kilograms", "kilogram", "kg", "tonnes", "tonne", "tons", "ton"),
        "length_in": ("centimeters", "centimeter", "cm", "millimeters", "millimeter", "mm", "meters", "meter", "metres", "metre", "feet", "foot", "ft"),
        "width_in": ("centimeters", "centimeter", "cm", "millimeters", "millimeter", "mm", "meters", "meter", "metres", "metre", "feet", "foot", "ft"),
        "height_in": ("centimeters", "centimeter", "cm", "millimeters", "millimeter", "mm", "meters", "meter", "metres", "metre", "feet", "foot", "ft"),
        "available_volume_cuft": ("cubicmeters", "cubicmeter", "m3", "liters", "liter", "litres", "litre", "gallons", "gallon", "liter", "litre"),
    }
    return any(name.endswith(unit) for unit in suffixes.get(field_name, ()))


def _suggest_mapping(headers: list[str], dataset_type: str) -> tuple[dict[str, str | None], list[str]]:
    fields = ORDER_FIELDS if dataset_type == "orders" else FLIGHT_FIELDS
    mapping: dict[str, str | None] = {}
    signals: list[str] = []
    used: set[str] = set()
    for internal in fields:
        aliases = ALIASES[internal]
        exact = next((header for header in headers if _norm(header) in aliases and header not in used and not _has_incompatible_unit(header, internal)), None)
        if exact is None:
            # Prefix matches support innocuous labels (e.g. "Order ID number"),
            # while incompatible unit suffixes are deliberately left unmapped.
            possible = [header for header in headers if _norm(header) not in used and any(_norm(header).startswith(alias) for alias in aliases) and not _has_incompatible_unit(header, internal) and not (internal == "flight_cost" and _norm(header).startswith("costperorder"))]
            exact = possible[0] if len(possible) == 1 else None
        mapping[internal] = exact
        if exact:
            used.add(exact)
            signals.append(f"{exact} looks like {internal}")
    return mapping, signals


def detect_type(
    headers: list[str],
    filename: str = "",
    rows: list[dict[str, str]] | None = None,
) -> tuple[str, str, list[str], dict[str, str | None]]:
    """Classify using mapped headers, sample values, and supporting filename hints.

    A filename by itself can never classify a file. At least three recognizable
    fields and a type-specific anchor are required before a filename/sample hint
    can resolve a close match.
    """
    order_map, order_signals = _suggest_mapping(headers, "orders")
    flight_map, flight_signals = _suggest_mapping(headers, "flight_capacity")
    order_score = sum(value is not None for value in order_map.values())
    flight_score = sum(value is not None for value in flight_map.values())
    order_anchor = bool(order_map.get("order_id") and (order_map.get("household_id") or order_map.get("product_name")))
    flight_anchor = bool(flight_map.get("departure_id") and flight_map.get("departure_date"))
    sample_rows = (rows or [])[:5]

    def sample_score(mapping: dict[str, str | None], field_names: tuple[str, ...]) -> int:
        return sum(
            1 for name in field_names
            if mapping.get(name) and any((row.get(mapping[name] or "") or "").strip() for row in sample_rows)
        )

    order_sample_score = sample_score(order_map, ("order_id", "household_id", "order_date", "destination_community"))
    flight_sample_score = sample_score(flight_map, ("departure_id", "departure_date", "available_totes", "available_payload_lb"))
    filename_hint = _norm(Path(filename).stem)
    order_filename_hint = any(token in filename_hint for token in ("order", "grocery", "household"))
    flight_filename_hint = any(token in filename_hint for token in ("flight", "capacity", "aircraft", "departure"))

    # Header mapping determines the type; sample and filename corroborate it.
    if order_anchor and order_score >= 4 and order_score > flight_score:
        confidence = "high" if order_score >= 7 and order_map.get("weight_lb") else "medium"
        signals = list(order_signals)
        if order_filename_hint:
            signals.append(f"filename {filename!r} suggests orders")
        if order_sample_score:
            signals.append(f"sample rows contain {order_sample_score} recognizable order fields")
        return "orders", confidence, signals, order_map
    if flight_anchor and flight_score >= 3 and flight_score > order_score:
        confidence = "high" if flight_map.get("available_totes") and flight_map.get("available_payload_lb") else "medium"
        signals = list(flight_signals)
        if flight_filename_hint:
            signals.append(f"filename {filename!r} suggests flight capacity")
        if flight_sample_score:
            signals.append(f"sample rows contain {flight_sample_score} recognizable capacity fields")
        return "flight_capacity", confidence, signals, flight_map

    # A close header match may use both hints, but cannot be decided by a name alone.
    if order_anchor and order_score >= 3 and order_sample_score >= 3 and order_filename_hint and order_score >= flight_score:
        return "orders", "medium", order_signals + ["filename and sample values support an orders dataset"], order_map
    if flight_anchor and flight_score >= 3 and flight_sample_score >= 3 and flight_filename_hint and flight_score >= order_score:
        return "flight_capacity", "medium", flight_signals + ["filename and sample values support a flight-capacity dataset"], flight_map
    return "unknown", "low", [], order_map if order_score >= flight_score else flight_map


def analyze_upload(filename: str, content: bytes) -> ImportDraft:
    name = safe_filename(filename)
    draft = ImportDraft(draft_id=secrets.token_urlsafe(18), filename=name)
    try:
        if Path(name).suffix.lower() == ".csv":
            draft.headers, draft.rows = _parse_csv(content, name)
        else:
            draft.headers, draft.rows, draft.note = _parse_xlsx(content, name)
        draft.dataset_type, draft.confidence, draft.signals, _mapping = detect_type(
            draft.headers, filename=draft.filename, rows=draft.rows
        )
    except ValueError as exc:
        draft.parse_error = str(exc)
    return draft


def _parse_number(value: str, field_name: str, row_number: int = 0, *, positive: bool = False) -> float:
    raw = value.strip() if isinstance(value, str) else ""
    if "," in raw and not re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", raw):
        raise ValueError(f"{field_name} must be numeric")
    try:
        number = float(raw.replace(",", ""))
    except (ValueError, AttributeError) as exc:
        raise ValueError(f"{field_name} must be numeric") from exc
    if not math.isfinite(number) or number < 0 or (positive and number <= 0):
        requirement = "a finite positive number" if positive else "a finite non-negative number"
        raise ValueError(f"{field_name} must be {requirement}")
    return number


def _parse_int(value: str, field_name: str, *, allow_prefix: bool = False) -> int:
    raw = value.strip()
    if "," in raw and not re.fullmatch(r"\d{1,3}(,\d{3})+", raw):
        raise ValueError(f"{field_name} must be an integer")
    raw = raw.replace(",", "")
    if allow_prefix:
        match = re.fullmatch(r"[A-Za-z_-]*(\d+)[A-Za-z_-]*", raw)
        if match:
            raw = match.group(1)
    try:
        number = float(raw)
    except ValueError as exc:
        raise ValueError(f"{field_name} must be an integer") from exc
    if not math.isfinite(number) or not number.is_integer() or number < 0:
        raise ValueError(f"{field_name} must be a non-negative integer")
    return int(number)


def _parse_date(value: str) -> date:
    raw = value.strip()
    try:
        return date.fromisoformat(raw)
    except ValueError:
        pass
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for pattern in ("%m/%d/%Y", "%Y/%m/%d", "%d-%b-%Y", "%b %d, %Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    raise ValueError("must be a date such as YYYY-MM-DD")


def preview_draft(
    draft: ImportDraft,
    dataset_type: str,
    mapping: dict[str, str | None],
    excluded_row_numbers: set[int] | None = None,
) -> dict[str, Any]:
    if dataset_type not in {"orders", "flight_capacity"}:
        return {"dataset_type": dataset_type, "valid": False, "error": "Select Orders or Flight / Capacity.", "mapping_errors": [], "valid_row_numbers": [], "invalid_rows": [], "valid_row_count": 0, "row_count": len(draft.rows), "column_count": len(draft.headers), "warnings": []}
    if draft.parse_error:
        return {"dataset_type": dataset_type, "valid": False, "error": draft.parse_error, "mapping_errors": [], "valid_row_numbers": [], "invalid_rows": [], "valid_row_count": 0, "row_count": 0, "column_count": len(draft.headers), "warnings": []}
    fields = ORDER_FIELDS if dataset_type == "orders" else FLIGHT_FIELDS
    required = ORDER_REQUIRED if dataset_type == "orders" else FLIGHT_REQUIRED
    mapping_errors: list[str] = []
    normalized_mapping: dict[str, str | None] = {}
    for internal in fields:
        column = mapping.get(internal)
        if column == "":
            column = None
        if column is not None and column not in draft.headers:
            mapping_errors.append(f"{fields[internal]} maps to missing column {column!r}.")
        if internal == "flight_cost" and column is not None and _norm(column).startswith("costperorder"):
            mapping_errors.append("Cost per order cannot be imported as flight cost; provide total flight cost instead.")
        if internal in required and not column:
            mapping_errors.append(f"Map a column to required field: {fields[internal]}.")
        normalized_mapping[internal] = column
    if len([column for column in normalized_mapping.values() if column]) != len(set(column for column in normalized_mapping.values() if column)):
        mapping_errors.append("A source column cannot be mapped to more than one field.")
    if mapping_errors:
        return {"dataset_type": dataset_type, "valid": False, "error": None, "mapping_errors": mapping_errors, "mapping": normalized_mapping, "valid_row_numbers": [], "invalid_rows": [], "valid_row_count": 0, "row_count": len(draft.rows), "column_count": len(draft.headers), "warnings": []}

    invalid: list[dict[str, Any]] = []
    valid_rows: list[int] = []
    seen_capacity_ids: dict[int, int] = {}
    order_metadata: dict[int, tuple[int, date, str, int]] = {}
    for row_number, row in enumerate(draft.rows, start=2):
        issues: list[str] = []
        def value(name: str) -> str:
            column = normalized_mapping.get(name)
            return (row.get(column or "") or "").strip()
        try:
            if dataset_type == "orders":
                order_id = _parse_int(value("order_id"), "Order ID")
                household_id = _parse_int(value("household_id"), "Household ID")
                order_date = _parse_date(value("order_date"))
                destination = value("destination_community")
                if not destination:
                    raise ValueError("Destination community is required")
                _parse_number(value("weight_lb"), "Weight")
                for name in ("length_in", "width_in", "height_in"):
                    _parse_number(value(name), ORDER_FIELDS[name], positive=True)
                batch_id = _parse_int(value("batch_id"), "Batch ID") if value("batch_id") else 1
                if value("quantity") and _parse_int(value("quantity"), "Quantity") <= 0:
                    raise ValueError("Quantity must be greater than zero")
                metadata = (household_id, order_date, destination, batch_id)
                if order_id in order_metadata and order_metadata[order_id] != metadata:
                    raise ValueError("Order ID has inconsistent household, date, destination, or batch metadata")
                order_metadata.setdefault(order_id, metadata)
            else:
                departure_id = _parse_int(value("departure_id"), "Flight ID", allow_prefix=True)
                _parse_date(value("departure_date"))
                totes = _parse_int(value("available_totes"), "Tote capacity")
                _parse_number(value("available_payload_lb"), "Payload capacity")
                volume_text = value("available_volume_cuft")
                if volume_text and volume_text.lower() not in {"unknown", "none", "null", "n/a", "na"}:
                    _parse_number(volume_text, "Volume capacity")
                cost_text = value("flight_cost")
                if cost_text and cost_text.lower() not in {"unknown", "none", "null", "n/a", "na"}:
                    _parse_number(cost_text, "Flight cost")
                currency_text = value("currency")
                if currency_text and not re.fullmatch(r"[A-Za-z]{3}", currency_text.strip()):
                    raise ValueError("Currency must be a three-letter code such as CAD")
                if departure_id in seen_capacity_ids:
                    raise ValueError(f"Duplicate flight/departure ID also appears on row {seen_capacity_ids[departure_id]}")
                seen_capacity_ids[departure_id] = row_number
        except ValueError as exc:
            issues.append(str(exc))
        if issues:
            invalid.append({"row_number": row_number, "issues": issues})
        else:
            valid_rows.append(row_number)

    if not draft.rows:
        return {"dataset_type": dataset_type, "valid": False, "error": "The file contains no data rows.", "mapping_errors": [], "valid_row_numbers": [], "invalid_rows": [], "valid_row_count": 0, "row_count": 0, "column_count": len(draft.headers), "warnings": []}

    excluded = excluded_row_numbers or set()
    unknown_exclusions = sorted(excluded - {item["row_number"] for item in invalid})
    if unknown_exclusions:
        invalid.append({"row_number": 0, "issues":[f"Cannot exclude rows not marked invalid: {unknown_exclusions}"]})
    remaining_invalid = [item for item in invalid if item["row_number"] not in excluded]
    warnings: list[str] = []
    optional_id = "product_id" if dataset_type == "orders" else None
    if dataset_type == "orders" and not normalized_mapping.get(optional_id or ""):
        warnings.append("Product IDs will be assigned from the source row number because no product ID column is mapped.")
    if dataset_type == "orders" and not normalized_mapping.get("product_name"):
        warnings.append("Product names will be generated from product IDs because no product/name column is mapped.")
    if dataset_type == "orders":
        warnings.append("Mapped weights are per-unit lb and dimensions are per-unit inches; quantities multiply the per-unit weight and box volume. Values are not automatically converted from other units.")
    if dataset_type == "flight_capacity":
        warnings.append("Mapped payload is interpreted as lb and cargo volume as ft³. Values are not automatically converted from other units.")
    if dataset_type == "flight_capacity" and not normalized_mapping.get("available_volume_cuft"):
        warnings.append("Cargo volume capacity is missing and will remain Unknown; no volume limit will be invented.")
    if draft.note:
        warnings.append(draft.note)
    return {
        "dataset_type": dataset_type,
        "valid": bool(valid_rows) and not remaining_invalid and not mapping_errors,
        "error": None,
        "mapping_errors": mapping_errors,
        "mapping": normalized_mapping,
        "valid_row_numbers": valid_rows,
        "invalid_rows": invalid,
        "valid_row_count": len(valid_rows),
        "row_count": len(draft.rows),
        "column_count": len(draft.headers),
        "headers": draft.headers,
        "sample_rows": draft.rows[:5],
        "excluded_row_numbers": sorted(excluded),
        "warnings": warnings,
    }


def normalize_orders(draft: ImportDraft, mapping: dict[str, str | None], excluded: set[int]) -> list[HouseholdOrder]:
    preview = preview_draft(draft, "orders", mapping, excluded)
    if not preview.get("valid"):
        raise ValueError("order dataset is not valid for import")
    valid_numbers = set(preview["valid_row_numbers"])
    by_order: dict[int, HouseholdOrder] = {}
    for row_number, row in enumerate(draft.rows, start=2):
        if row_number not in valid_numbers:
            continue
        def value(name: str) -> str:
            column = mapping.get(name)
            return (row.get(column or "") or "").strip()
        order_id = _parse_int(value("order_id"), "Order ID")
        household_id = _parse_int(value("household_id"), "Household ID")
        order_date = _parse_date(value("order_date"))
        destination = value("destination_community")
        batch_id = _parse_int(value("batch_id"), "Batch ID") if value("batch_id") else 1
        raw_product_id = value("product_id")
        # Preserve the source identifier exactly, including alphanumeric SKUs
        # and leading zeroes. The deterministic row fallback is used only when
        # the source has no product-ID column/value.
        product_id: int | str = raw_product_id if raw_product_id else row_number
        product_name = value("product_name") or f"Product {product_id}"
        quantity = _parse_int(value("quantity"), "Quantity") if value("quantity") else 1
        item = OrderItem(
            order_id=order_id, household_id=household_id, batch_id=batch_id,
            order_date=order_date, destination_community=destination,
            product_id=product_id, product_name=product_name,
            length_in=_parse_number(value("length_in"), "Length", positive=True),
            width_in=_parse_number(value("width_in"), "Width", positive=True),
            height_in=_parse_number(value("height_in"), "Height", positive=True),
            weight_lb=_parse_number(value("weight_lb"), "Weight"),
            order_item_id=f"user:{draft.draft_id}:row:{row_number}:order:{order_id}",
            source_quantity=quantity,
        )
        if order_id not in by_order:
            by_order[order_id] = HouseholdOrder(
                order_id=order_id, household_id=household_id, batch_id=batch_id,
                order_date=order_date, destination_community=destination, items=[item],
            )
        else:
            by_order[order_id].items.append(item)
    return list(by_order.values())


def normalize_capacities(draft: ImportDraft, mapping: dict[str, str | None], excluded: set[int]) -> list[FlightCapacity]:
    preview = preview_draft(draft, "flight_capacity", mapping, excluded)
    if not preview.get("valid"):
        raise ValueError("flight capacity dataset is not valid for import")
    valid_numbers = set(preview["valid_row_numbers"])
    result: list[FlightCapacity] = []
    for row_number, row in enumerate(draft.rows, start=2):
        if row_number not in valid_numbers:
            continue
        def value(name: str) -> str:
            column = mapping.get(name)
            return (row.get(column or "") or "").strip()
        volume_text = value("available_volume_cuft")
        volume = None if not volume_text or volume_text.lower() in {"unknown", "none", "null", "n/a", "na"} else _parse_number(volume_text, "Volume capacity")
        cost_text = value("flight_cost")
        cost = None if not cost_text or cost_text.lower() in {"unknown", "none", "null", "n/a", "na"} else _parse_number(cost_text, "Flight cost")
        currency = value("currency").strip().upper() or ("CAD" if cost is not None else None)
        result.append(FlightCapacity(
            departure_id=_parse_int(value("departure_id"), "Flight ID", allow_prefix=True),
            departure_date=_parse_date(value("departure_date")),
            available_totes=_parse_int(value("available_totes"), "Tote capacity"),
            available_payload_lb=_parse_number(value("available_payload_lb"), "Payload capacity"),
            available_volume_cuft=volume,
            flight_cost=cost,
            currency=currency,
        ))
    return result
